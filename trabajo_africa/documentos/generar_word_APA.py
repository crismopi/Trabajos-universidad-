from pathlib import Path
import re,json,shutil
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
p=Path(__file__).resolve().parent.parent
refs=json.loads((p/'datos/revisiones/referencias_informe.json').read_text()); R={r['id']:r for r in refs}
# Explicit metadata normalization; no invented author for the collective 2006 report.
auth={1:'Gamba, P., & Kibaara, B.',2:'Kenya National Bureau of Statistics',3:'National Council for Population and Development',4:'Central Bank of Kenya',5:'Kenya National Bureau of Statistics',6:'Shah, S., Zamara, & Centre for Affordable Housing Finance in Africa',7:'InfoTradeKenya',8:'East African Community',9:'Banco Mundial',10:'International Finance Corporation',11:'UNCTAD',12:'UNCTAD',13:'Private Infrastructure Development Group',14:'Constitution of Kenya',15:'',16:'Banco Mundial',17:'The National Treasury',18:'Banco Mundial',19:'Acemoglu, D., & Robinson, J. A.',20:'Freedom House',21:'Centric Africa Ltd.',22:'Acorn Holdings',23:'Robinson, H. J., Mwangi, B., & Atkinson, B.',24:'Nairobi City County Government'}
year={i:re.search(r'\d{4}',r['fecha_publicacion'])[0] if re.search(r'\d{4}',r['fecha_publicacion']) else 's. f.' for i,r in R.items()};year[10]='2024';year[11]='2025b';year[12]='2025a'
short={i:a for i,a in auth.items()};short.update({1:'Gamba & Kibaara',6:'Shah et al.',14:'Constitution of Kenya',15:'*Kenya National Human Development Report 2006*',19:'Acemoglu & Robinson',23:'Robinson et al.'})
titles={i:r['titulo'] for i,r in R.items()};titles.update({1:'Structural implications of economic liberalization on agriculture and rural development: RuralStruc program, phase I',3:'Press release',4:'Annual GDP',6:'Affordable housing access for Kenya’s formal workforce: What pension data reveals and what can work',7:'AfCFTA market requirements',8:'EAC at a glance',9:'Kenya systematic country diagnostic',10:'IHS Kenya: Environmental and social review summary (Project 47136)',11:'World investment report 2025: Country fact sheet—Kenya',12:'COMESA investment report 2025: Investment trends and policy insights',13:'PIDG anchors Acorn’s new REIT to finance affordable housing for Kenya’s young urbanites',15:'Kenya national human development report 2006: Human security and human development—A deliberate choice',16:'Reforms are needed to stem a slow-down in Kenya’s economic growth',17:'Budget review and outlook paper',18:'Kenya economic update 32: From barriers to bridges—Procompetitive reforms for productivity and jobs in Kenya',19:'Why nations fail: The origins of power, prosperity, and poverty',20:'Freedom in the world 2026: Kenya',21:'Environmental and social impact assessment study report for the proposed Riruta Keza IHS housing developments',22:'Overview',23:'Understanding the tenants of Fusion Capital’s rental properties in Nairobi, Kenya',24:'Customer service charter'})
def sortkey(i):return (auth[i] or titles[i]).lower(),year[i],titles[i].lower()
def citation(m):
 ids=sorted(set(map(int,re.findall(r'\d+',m[0]))),key=sortkey)
 return '('+'; '.join(short[i]+', '+year[i] for i in ids)+')'
f=p/'borrador_informe_kenia.md';original=f.read_text();backup=p/'datos/revisiones/borrador_antes_APA_2026-10-21.md'
if not backup.exists():backup.write_text(original)
s=backup.read_text().split('## Bibliografía y fuentes')[0];s=re.sub(r'\[\d+\](?:, \[\d+\])*',citation,s)
s=s.replace('**Figura 1. Crecimiento de Kenia: crisis y recuperación**','**Figura 1**\n\n*Crecimiento de Kenia: crisis y recuperación*')
s=s.replace('*Fuente: elaboración propia con KNBS, Economic Survey 2026 (Kenya National Bureau of Statistics, 2026), y CBK, Annual GDP (Central Bank of Kenya, s. f.). Años seleccionados; serie revisada.*','*Nota.* Elaboración propia a partir de *Economic Survey 2026* (Kenya National Bureau of Statistics, 2026) y *Annual GDP* (Central Bank of Kenya, s. f.). Años seleccionados; serie revisada.')
# Initial paragraph: place historical citation on relevant statement.
s=s.replace('a una mayor participación privada.','a una mayor participación privada (Gamba & Kibaara, 2007).',1).replace('Se elige Acorn Holdings como empresa de cabecera y se compara con IHS y Fusion Capital (Gamba & Kibaara, 2007).','Se elige Acorn Holdings como empresa de cabecera y se compara con IHS y Fusion Capital.')
s=s.replace('(UNCTAD, 2025a; UNCTAD, 2025b)', '(UNCTAD, 2025a, 2025b)')
s=s.replace('(graficos/01_crecimiento_pib.svg)', '(graficos/01_crecimiento_pib_APA.svg)')
s=s.replace('El informe de desarrollo humano de Kenia del PNUD describe','El informe de desarrollo humano de Kenia patrocinado por PNUD describe')
dates={3:'2025, 12 de marzo',6:'2026, 19 de mayo',13:'2025, 17 de noviembre',16:'2024, 10 de diciembre',23:'2025, 24 de febrero',24:'2025, junio'}
entries=[]
for i in sorted(R,key=sortkey):
 date=dates.get(i,year[i]);title=titles[i];a=auth[i];url=R[i]['url'];pub={1:'Egerton University, Tegemeo Institute.',6:'FSD Kenya.',19:'Crown.',21:'National Environment Management Authority.',23:'Reall.'}.get(i,'')
 if i==14:entry=f'*Constitution of Kenya*, art. 65 (2010). Kenya Law Reform Commission. {url}'
 elif not a:entry=f'*{title}*. ({date}). {url}'
 else:entry=f'{a} ({date}). *{title}*. '+(pub+' ' if pub else '')+('Recuperado el 21 de octubre de 2026, de ' if i in [4,7,8,22] else '')+url
 entries.append((i,entry))
s+='## Referencias\n\n'+'\n\n'.join(e for _,e in entries)+'\n';f.write_text(s)
(p/'datos/revisiones/referencias_APA.json').write_text(json.dumps([{'id_anterior':i,'autor':auth[i],'año_cita':year[i],'titulo':titles[i],'referencia':e,'url':R[i]['url']} for i,e in entries],ensure_ascii=False,indent=2)+'\n')
# Clean chart artwork: APA figure number/title/note are in the document.
svg=p/'graficos/01_crecimiento_pib.svg';v=svg.read_text();v=re.sub(r'<text x="55" y="60".*?</text>','',v);v=re.sub(r'<text x="55" y="625".*?</text>','',v);v=re.sub(r'<text x="55" y="648".*?</text>','',v);v=v.replace('height="680" viewBox="0 0 1200 680"','height="520" viewBox="0 75 1200 520"');(p/'graficos/01_crecimiento_pib_APA.svg').write_text(v)
# Word layout: APA citation style, restrained academic typography.
d=Document();sec=d.sections[0];sec.top_margin=sec.bottom_margin=Inches(.85);sec.left_margin=sec.right_margin=Inches(.9)
sec.page_width=Inches(8.27);sec.page_height=Inches(11.69)
styles=d.styles
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
 styles[name].font.name='Arial';styles[name].font.color.rgb=RGBColor(0,0,0)
 for bd in list(styles[name].element.xpath('./w:pPr/w:pBdr')):bd.getparent().remove(bd)
styles['Normal'].font.size=Pt(11);pf=styles['Normal'].paragraph_format;pf.line_spacing=1.4;pf.space_after=Pt(6)
for name,size in [('Title',20),('Heading 1',14),('Heading 2',12),('Heading 3',11)]:styles[name].font.size=Pt(size);styles[name].paragraph_format.keep_with_next=True
# page number
foot=sec.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.RIGHT;fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');foot._p.append(fld)
d.core_properties.author='Cristina Moro';d.core_properties.title='Kenia y la vivienda de alquiler para trabajadores en Nairobi'
def rich(pa,t):
 for token in re.split(r'(\*\*.*?\*\*|\*[^*]+\*)',t):
  if not token:continue
  r=pa.add_run(token.strip('*') if token.startswith('*') else token)
  if token.startswith('**'):r.bold=True
  elif token.startswith('*'):r.italic=True
 return pa
lines=s.splitlines();i=0;inrefs=False
while i<len(lines):
 l=lines[i].strip();i+=1
 if not l:continue
 if l.startswith('# '):d.add_paragraph('Kenia y la vivienda de alquiler para trabajadores en Nairobi','Title');continue
 if l.startswith('## Referencias'):d.add_page_break();d.add_heading('Referencias',level=1);inrefs=True;continue
 if l.startswith('## '):d.add_heading(l[3:].replace(':',''),1);continue
 if l.startswith('### '):d.add_heading(l[4:].replace(':',''),2);continue
 if l.startswith('|'):
  rows=[l]
  while i<len(lines) and lines[i].strip().startswith('|'):rows.append(lines[i].strip());i+=1
  cells=[[c.strip() for c in r.strip('|').split('|')] for r in rows if not re.match(r'^\|[-| :]+\|$',r)]
  table=d.add_table(rows=0,cols=len(cells[0]));table.alignment=WD_TABLE_ALIGNMENT.CENTER;table.autofit=False
  widths=[3.235,3.235] if len(cells[0])==2 else [1.05,1.8,1.3,.95,1.37]
  for col,w in zip(table.columns,widths):col.width=Inches(w)
  for ri,row in enumerate(cells):
   cs=table.add_row().cells
   trPr=table.rows[-1]._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');trPr.append(cant)
   if ri==0:trPr.append(OxmlElement('w:tblHeader'))
   for j,(c,t) in enumerate(zip(cs,row)):
    c.width=Inches(widths[j]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    pr=c._tc.get_or_add_tcPr();borders=OxmlElement('w:tcBorders')
    for edge in ['top','left','bottom','right']:
     e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
    pr.append(borders);shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E8EEF1' if ri==0 else ('F6F8F9' if ri%2==0 else 'FFFFFF'));pr.append(shade)
    mar=OxmlElement('w:tcMar')
    for edge in ['top','left','bottom','right']:
     e=OxmlElement('w:'+edge);e.set(qn('w:w'),'90');e.set(qn('w:type'),'dxa');mar.append(e)
    pr.append(mar);pa=rich(c.paragraphs[0],t);pa.paragraph_format.line_spacing=1.1;pa.paragraph_format.space_after=Pt(3);pa.paragraph_format.space_before=Pt(3)
    for r in pa.runs:r.font.size=Pt(9 if len(cs)>2 else 10);r.bold=ri==0 or r.bold
  d.add_paragraph().paragraph_format.space_after=Pt(0);continue
 if l.startswith('!['):
  pa=d.add_paragraph();pa.paragraph_format.keep_with_next=True;pa.add_run().add_picture(str(p/'graficos/01_crecimiento_pib_APA.png'),width=Inches(6.4));continue
 pa=rich(d.add_paragraph(),l)
 if inrefs:pa.paragraph_format.left_indent=Inches(.5);pa.paragraph_format.first_line_indent=Inches(-.5);pa.paragraph_format.line_spacing=1.2;pa.paragraph_format.keep_together=True;[setattr(r.font,'size',Pt(10.5)) for r in pa.runs]
 if l.startswith('**Figura') or l.startswith('*Crecimiento'):pa.paragraph_format.keep_with_next=True
 if l.startswith('*Nota.'):pa.paragraph_format.line_spacing=1.1;pa.paragraph_format.space_after=Pt(10);[setattr(r.font,'size',Pt(9)) for r in pa.runs]
# Save only after APA figure PNG has been generated.
def save():d.save(p/'documentos/Trabajo_Kenia_Cristina_Moro_APA.docx')
if __name__=='__main__':save()
